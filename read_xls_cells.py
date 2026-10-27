#!/usr/bin/env python3
"""Dump the first sheet of a Pub 1304 .xls (or .xlsx) as CSV on stdout, one
record per non-empty cell plus one per merged range:

    C,row,col,flag,value      a cell (1-based); value as text, numbers as shortest round-trip repr
    M,first_row,last_row,first_col,last_col   a merged range (1-based inclusive),
                              or a header centred across a run of cells (below)

Used by alignment_helpers.R::read_bysize_sheet(). It exists because SOI marks
its disclosure flags with NUMBER FORMATS, not text: a cell shown as "* 1,234"
holds 1234 under the format '"* "#,##0', and a cell shown as "**" holds 0
under '"** "#,##0;...'. readxl returns the bare numbers and loses both marks,
so a combined cell reads as a true zero. The flag column carries them:

    caution    format prefix "* "   estimate flagged for sampling variability
    combined   format prefix "** "  combined with another cell to prevent
                                    disclosure; a zero here is not a zero
    ''         anything else

Requires xlrd. It reads both formats in the store: BIFF8 (TY2004 on) and
the BIFF4 of Table 2.3's TY1996-2003 files (TY1997 aside). The flag formats
begin in TY2005; earlier vintages carry none (checked 2026-09-26), so an
empty flag column there is what SOI published, not a reader gap.

A .xlsx file (Table 4.1 from TY2023) is read with openpyxl instead, which
xlrd 2 no longer opens. Its records are the same: the flag comes from the
same number formats, merged ranges from the sheet's merge list, and a header
centred across a selection from the "centerContinuous" alignment.
"""

import csv
import re
import sys

import xlrd

if len(sys.argv) != 2:
    sys.exit("usage: read_xls_cells.py <file.xls>")

out = csv.writer(sys.stdout, lineterminator="\n")

# The flag is decided by the positive section of the format, which every
# flagged format in the store opens with a quoted "* " or "** ".
FLAG_RE = re.compile(r'^"(\*+)\s*"')


def flag_of(format_str):
    m = FLAG_RE.match(format_str or "")
    if not m:
        return ""
    return "combined" if len(m.group(1)) >= 2 else "caution"


if sys.argv[1].lower().endswith(".xlsx"):
    import openpyxl

    ws = openpyxl.load_workbook(sys.argv[1]).worksheets[0]
    for row in ws.iter_rows():
        for cell in row:
            v = cell.value
            if v is None:
                continue
            if isinstance(v, bool) or not isinstance(v, (int, float)):
                text = str(v)
                if text.strip() == "":
                    continue
                flag = ""
            else:
                # float() so an integer prints as xlrd would ("12.0")
                text = repr(float(v))
                flag = flag_of(cell.number_format)
            out.writerow(["C", cell.row, cell.column, flag, text])
    for rng in sorted(ws.merged_cells.ranges, key=lambda g: (g.min_row, g.min_col)):
        out.writerow(["M", rng.min_row, rng.max_row, rng.min_col, rng.max_col])
    for row in ws.iter_rows():
        c = 0
        while c < len(row):
            cell = row[c]
            if cell.value is None or str(cell.value).strip() == "" or \
                    cell.alignment.horizontal != "centerContinuous":
                c += 1
                continue
            end = c + 1
            while end < len(row) and (row[end].value is None or str(row[end].value).strip() == "") \
                    and row[end].alignment.horizontal == "centerContinuous":
                end += 1
            if end - c >= 2:
                out.writerow(["M", cell.row, cell.row, c + 1, end])
            c = end
    sys.exit(0)

book = xlrd.open_workbook(sys.argv[1], formatting_info=True)
sheet = book.sheet_by_index(0)


def cell_flag(r, c):
    fmt = book.format_map[book.xf_list[sheet.cell_xf_index(r, c)].format_key]
    return flag_of(fmt.format_str)


for r in range(sheet.nrows):
    for c in range(sheet.ncols):
        ctype = sheet.cell_type(r, c)
        if ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
            continue
        v = sheet.cell_value(r, c)
        if ctype == xlrd.XL_CELL_NUMBER:
            text = repr(v)
            flag = cell_flag(r, c)
        else:
            text = str(v)
            if text.strip() == "":
                continue
            flag = ""
        out.writerow(["C", r + 1, c + 1, flag, text])

for r0, r1, c0, c1 in sorted(sheet.merged_cells):
    # xlrd ranges are 0-based half-open
    out.writerow(["M", r0 + 1, r1, c0 + 1, c1])

# BIFF4 has no merged cells: those vintages (Table 2.3 TY1996-2003) span a
# header over its columns with "centre across selection" instead -- the text
# in the leftmost cell, the empty cells to its right carrying the same
# alignment. Each such run is emitted as a one-row merged range so the R side
# treats it as a spanner.
CENTRE_ACROSS = 6


def is_empty(r, c):
    return (sheet.cell_type(r, c) in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK)
            or str(sheet.cell_value(r, c)).strip() == "")


def centred_across(r, c):
    return book.xf_list[sheet.cell_xf_index(r, c)].alignment.hor_align == CENTRE_ACROSS


for r in range(sheet.nrows):
    c = 0
    while c < sheet.ncols:
        if is_empty(r, c) or not centred_across(r, c):
            c += 1
            continue
        end = c + 1
        while end < sheet.ncols and is_empty(r, end) and centred_across(r, end):
            end += 1
        if end - c >= 2:
            out.writerow(["M", r + 1, r + 1, c + 1, end])
        c = end
